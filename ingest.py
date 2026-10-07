import io, re, zipfile, requests
from urllib.parse import urlparse, parse_qs
from bs4 import BeautifulSoup

IMG = {'png': 'image/png', 'jpg': 'image/jpeg', 'jpeg': 'image/jpeg'}
TEXTUAL = {'txt', 'md', 'markdown', 'csv', 'json', 'xml'}

def _yt_id(url):
    u = urlparse(url)
    if u.hostname and u.hostname.endswith('youtu.be'): return u.path.strip('/')
    return parse_qs(u.query).get('v', [u.path.split('/')[-1]])[0]

def _html_text(markup):
    s = BeautifulSoup(markup, 'html.parser')
    for t in s(['script', 'style', 'nav', 'footer', 'header', 'aside', 'noscript', 'iframe']): t.decompose()
    return s, (s.find('article') or s.find('main') or s.body or s).get_text('\n')

def extract(file=None, url=None, text=None, title=None, ocr=None):
    """returns (name, type, [(page, text)])"""
    if file:
        name, data = file
        ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
        if ext == 'pdf':
            from pypdf import PdfReader
            return name, 'pdf', [(i + 1, p.extract_text() or '') for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)]
        if ext == 'docx':
            from docx import Document
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows: parts.append(' | '.join(c.text for c in r.cells))
            return name, 'docx', [(None, '\n'.join(parts))]
        if ext == 'pptx':
            from pptx import Presentation
            pages = []
            for i, sl in enumerate(Presentation(io.BytesIO(data)).slides):
                t = '\n'.join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
                if t.strip(): pages.append((i + 1, t))
            return name, 'pptx', pages
        if ext == 'xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            pages = []
            for i, ws in enumerate(wb.worksheets):
                rows = [' | '.join('' if c is None else str(c) for c in r) for r in ws.iter_rows(values_only=True)]
                t = '\n'.join(r for r in rows if r.strip(' |'))
                if t.strip(): pages.append((i + 1, f'Sheet: {ws.title}\n{t}'))
            return name, 'xlsx', pages
    if url:
        if re.search(r'(youtube\.com|youtu\.be)', url):
            from youtube_transcript_api import YouTubeTranscriptApi
            vid = _yt_id(url)
            try: tr = YouTubeTranscriptApi().fetch(vid); t = ' '.join(s.text for s in tr)
            except AttributeError: t = ' '.join(s['text'] for s in YouTubeTranscriptApi.get_transcript(vid))
            return title or url, 'youtube', [(None, t)]
        r = requests.get(url, headers={'User-Agent': 'Mozilla/5.0 NotebookClone'}, timeout=20)
        r.raise_for_status()
        s, t = _html_text(r.text)
        return title or (s.title.string.strip() if s.title and s.title.string else url), 'web', [(None, t)]
    if text: return title or 'Pasted text', 'text', [(None, text)]
    raise ValueError('Nothing to add')

def chunk(pages, size=1200, overlap=150):
    """returns (chunks[(page, text, start)], normalized_pages[{page, text}])"""
    out, norm = [], []
    for page, text in pages:
        t = re.sub(r'\n{3,}', '\n\n', re.sub(r'[ \t]+', ' ', text)).strip()
        norm.append({'page': page or 0, 'text': t})
        i = 0
        while i < len(t):
            end = min(i + size, len(t))
            if end < len(t):
                w = t[end - 300:end]; cut = max(w.rfind('\n'), w.rfind('. '))
                if cut > 0: end = end - 300 + cut + 1
            raw = t[i:end]; piece = raw.strip()
            if len(piece) > 20: out.append((page, piece, i + len(raw) - len(raw.lstrip())))
            if end >= len(t): break
            i = max(end - overlap, i + 1)
    return out, norm
